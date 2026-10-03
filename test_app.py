import unittest
import io
import zipfile
import openpyxl
from app import app

class SummaryAppTestCase(unittest.TestCase):
    def setUp(self):
        self.app = app.test_client()
        self.app.testing = True

    def test_index_route(self):
        response = self.app.get('/')
        self.assertEqual(response.status_code, 200)
        self.assertIn(b'MSG Oilfield', response.data)

    def test_suppliers_api(self):
        response = self.app.get('/api/suppliers')
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertTrue(data['success'])
        self.assertIn("K.HASHIM LLC", data['suppliers'])

    def test_download_internal_stores_excel(self):
        payload = {
            "client": "Dubai Petroleum Establishment",
            "po_number": "50307533",
            "msg_ref": "MSG-1026-2171",
            "is_internal": True,
            "filter_machined_only": True,
            "items": [
                {
                    "sl_no": "1",
                    "client_po_item_no": "10000",
                    "description": "UNION 1\", 6000#",
                    "qty": "3 OF 5",
                    "uom": "pcs",
                    "heat_number": "HN123",
                    "certificate_number": "MTC-001",
                    "make": "WMAASS",
                    "remarks": "Machined from S.40 to Sch.20",
                    "supplier_name": "K.HASHIM LLC",
                    "supplier_po": "PO-1026-3789"
                }
            ]
        }
        response = self.app.post('/api/download', json=payload)
        self.assertEqual(response.status_code, 200)
        
        excel_bytes = io.BytesIO(response.data)
        wb = openpyxl.load_workbook(excel_bytes)
        ws = wb.active
        self.assertEqual(ws['A1'].value, "MSG OILFIELD - STORES SUMMARY SHEET")
        self.assertEqual(ws['A5'].value, "SL NO")
        self.assertEqual(ws['B5'].value, "CLIENT PO ITEM NO")
        self.assertEqual(ws['C5'].value, "Description")
        self.assertEqual(ws['B6'].value, "10000")

    def test_export_supplier_pdfs_zip(self):
        payload = {
            "client": "Dubai Petroleum Establishment",
            "po_number": "50307533",
            "msg_ref": "MSG-1026-2171",
            "items": [
                {
                    "sl_no": "1",
                    "client_po_item_no": "10000",
                    "description": "UNION 1\", 6000#",
                    "qty": "12.00",
                    "uom": "pcs",
                    "heat_number": "HN999",
                    "certificate_number": "CERT-111",
                    "make": "WMAASS",
                    "remarks": "Locally made from plate",
                    "supplier_name": "K.HASHIM LLC",
                    "supplier_po": "PO-1026-3789"
                }
            ]
        }
        response = self.app.post('/api/export-supplier-pdfs', json=payload)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, 'application/zip')

    def test_stock_supplier_and_offered_remark_parsing(self):
        from pdf_parser import strip_internal_product_code, strip_leading_item_number, REMARK_PATTERN
        from excel_generator import is_meaningful_item

        # 1. Test code with STOCK
        desc = "BL25150RFA105/A350LF2ST&H-STOCK FLANGE 2\", 150#, BLIND RF"
        cleaned = strip_internal_product_code(desc)
        self.assertEqual(cleaned, "FLANGE 2\", 150#, BLIND RF")

        # 2. Test leading item number stripping (e.g. '1 ECC. RED...' -> 'ECC. RED...')
        self.assertEqual(strip_leading_item_number("1 ECC. RED. 6\" X 3\"", sl_no=1), "ECC. RED. 6\" X 3\"")
        self.assertEqual(strip_leading_item_number("2 ECC. RED. 8\" X 6\"", sl_no=2), "ECC. RED. 8\" X 6\"")
        self.assertEqual(strip_leading_item_number("1\" FLANGE 150#", sl_no=1), "1\" FLANGE 150#")

        # 3. Test Offered remark matching
        item_offered = {
            "description": "UNION 1\", 6000#",
            "remarks": "Offered 1/2\" NPT connection"
        }
        self.assertTrue(is_meaningful_item(item_offered))

        # 4. Test remark pattern matching
        m = REMARK_PATTERN.search("SPECTACLE BLIND 3/4\" - Offered 1/2\" NPT")
        self.assertIsNotNone(m)
        self.assertEqual(m.group(0), "Offered 1/2\" NPT")

if __name__ == '__main__':
    unittest.main()
